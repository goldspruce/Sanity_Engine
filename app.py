"""Local browser UI:  pip install streamlit python-docx ;  streamlit run app.py
Paste text, upload .docx/.txt/.md, or give a Google Doc link shared as 'anyone with the link can view'."""
import io, json, os, re, tempfile, time, urllib.request
import streamlit as st
import engine, triage
from run import backend, books_default


def docx_text(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n\n".join(parts)


def gdoc_text(url):
    m = re.search(r"/document/d/([\w-]+)", url)
    if not m:
        raise ValueError("Not a Google Doc link")
    with urllib.request.urlopen(f"https://docs.google.com/document/d/{m.group(1)}/export?format=txt", timeout=30) as r:
        return r.read().decode("utf-8", "ignore")


def safe_json_dumps(obj, indent=1):
    def default(o):
        if isinstance(o, (set, tuple)):
            return list(o)
        return str(o)
    return json.dumps(obj, indent=indent, default=default)


st.set_page_config(page_title="Adversarial hypothesis checker", page_icon="🔬")
st.title("Adversarial hypothesis checker")

nb_url = os.environ.get("NOTEBOOKLM_URL", "").strip() or "https://notebooklm.google.com"
st.markdown(f"> To investigate untested claims or examine biological/physical mechanisms further, you can cross-examine here is a [link to a Google NotebookLM pre-loaded with these canonical textbooks]({nb_url})")

with st.expander("📚 Pre-Loaded Canonical Textbooks in Corpus & NotebookLM (11)", expanded=False):
    st.markdown(f"🔗 **Google NotebookLM Workspace:** [{nb_url}]({nb_url})")
    st.markdown("""
    All reference textbooks are indexed in `corpus/` for deterministic BM25 retrieval and quote verification:
    - **Alberts et al. (7th ed., 2022)**: `alberts_molecular_biology_of_the_cell_7ed_2022` *(Cell biology, cross-bridges, motors)*
    - **Berg (1993)**: `berg_random_walks_in_biology_1993` *(Diffusion, Brownian motion, motility, low Reynolds number)*
    - **Cover & Thomas (2nd ed., 2006)**: `cover_elements_of_information_theory_2ed_2006` *(Information theory, channel capacity, rate distortion)*
    - **Feynman et al. (Vol. 1, 1963)**: `feynman_lectures_on_physics_vol1_1963` *(Mechanics, oscillations, radiation, heat, acoustics)*
    - **Feynman et al. (Vol. 2, 1964)**: `feynman_lectures_on_physics_vol2_1964` *(Electromagnetism, wave propagation, dielectric media)*
    - **Johnston & Wu (1st ed., 1995)**: `johnston_foundations_of_cellular_neurophysiology_1ed_1995` *(Cellular neurophysiology, cable theory, ion channels)*
    - **Kreyszig (10th ed., 2011)**: `kreyszig_advanced_engineering_mathematics_10ed_2011` *(Wave equations, PDEs, Fourier/Laplace transforms)*
    - **Murray (2nd ed., 1993)**: `murray_mathematical_biology_an_introduction_2ed_1993` *(Mathematical biology, reaction-diffusion, population dynamics)*
    - **Nelson & Cox (8th ed., 2021)**: `nelson_lehninger_principles_of_biochemistry_8ed_2021` *(Biochemistry, ATP kinetics, bioenergetics, metabolic flux)*
    - **Strogatz (3rd ed., 2024)**: `strogatz_nonlinear_dynamics_and_chaos_3ed_2024` *(Nonlinear dynamics, limit cycles, bifurcations, oscillators)*
    - **Winter (5th ed., 2020)**: `winter_biomechanics_and_motor_control_of_human_movement_5ed_2020` *(Joint kinetics, EMG, muscle power, stretch-shortening cycle)*
    """)

with st.sidebar:
    gen_name = st.selectbox("Generator model", ["gemini", "anthropic", "mock"])
    critic_name = st.selectbox("Reviewer model", ["same as generator", "gemini", "anthropic", "mock"])
    corpus = st.text_input("Corpus folder", "corpus")
    if os.path.isdir(corpus):
        with st.expander("📚 Books in Corpus"):
            for bf in sorted(os.listdir(corpus)):
                if bf.lower().endswith((".txt", ".md")):
                    st.caption(f"• `{os.path.splitext(bf)[0]}`")
    books = st.text_input("Book names to focus on in model recall?", "",
                          help="Optional. Leave blank to default to all books in your corpus folder.")
    review = st.checkbox("Create expert triage cards (optional)")
    entail = st.checkbox("Auto entailment pre-check")
    check_math = st.checkbox("Math / number verification (optional)", value=False)

    if st.button("🔄 Reset / New Hypothesis", help="Clear current session and start a new hypothesis"):
        st.session_state["report"] = None
        st.session_state["latest_report"] = None
        st.session_state["followup_reports"] = []
        st.session_state["follow_up_history"] = []
        st.session_state["session_active"] = True
        st.rerun()

    st.divider()
    if st.button("🛑 End Session (Exit)", help="Stop the local server and exit Python process"):
        st.session_state["session_active"] = False
        st.warning("Session ended. Server stopped. You can safely close this browser tab.")
        time.sleep(0.5)
        os._exit(0)

if not st.session_state.get("session_active", True):
    st.warning("🛑 Session has ended. You can review past results below, or restart a new session.")
    if st.button("🔄 Start New Session", type="primary"):
        st.session_state["session_active"] = True
        st.session_state["report"] = None
        st.session_state["latest_report"] = None
        st.session_state["followup_reports"] = []
        st.session_state["follow_up_history"] = []
        st.rerun()

text = st.text_area("Hypothesis", height=200)
up = st.file_uploader("...or upload .docx / .txt / .md", type=["docx", "txt", "md"])
link = st.text_input("...or Google Doc link (must be viewable by anyone with the link)")
if up:
    text = docx_text(up.getvalue()) if up.name.lower().endswith(".docx") else up.getvalue().decode("utf-8", "ignore")
elif link and not text:
    try:
        text = gdoc_text(link)
    except Exception as e:
        st.error(f"Could not read the Google Doc: {e}")
if text:
    text = st.text_area("Text that will be checked (edit if extraction dropped anything)", text, height=200, key="confirm")

if st.button("Run", disabled=not text, type="primary"):
    try:
        gen = backend(gen_name)
        crit = gen if critic_name == "same as generator" else backend(critic_name)
        with st.spinner("Running adversarial review..."):
            rep = engine.run(text, books or books_default(corpus), gen, crit, engine.LocalBM25(corpus), auto_entail=entail,
                             check_math=check_math, log=lambda m: None)
            st.session_state["report"] = rep
            st.session_state["latest_report"] = rep
            st.session_state["followup_reports"] = []
            st.session_state["follow_up_history"] = []
            st.session_state["session_active"] = True
            if review:
                d = tempfile.mkdtemp()
                triage.build(rep, d)
                st.session_state["triage_files"] = {
                    fn: open(f"{d}/{fn}").read() for fn in ("triage.md", "responses.csv", "triage.json")
                }
            else:
                st.session_state["triage_files"] = None
    except Exception as e:
        st.error(f"{type(e).__name__}: {e}")

if st.session_state.get("report"):
    rep = st.session_state["report"]
    counts = rep.get("counts") or engine.count_flags(rep.get("issues", []), rep.get("green_flags", []))
    col_r, col_y, col_g = st.columns(3)
    col_r.metric("🚩 Red Flags", counts.get("red", 0), help="Fatal physical or biological invariant violations")
    col_y.metric("🟨 Yellow Flags", counts.get("yellow", 0), help="Empirical friction & parameter regime constraints (requires edge conditions)")
    col_g.metric("🟩 Green Flags", counts.get("green", 0), help="Biophysically sound mechanisms & canonical principles")

    rendered = engine.render(rep)
    st.markdown(rendered)
    st.divider()
    col1, col2 = st.columns(2)
    with col1:
        st.download_button("📥 Download report.md (Base Hypothesis)", rendered, "report.md", mime="text/markdown", key="dl_report_md")
    with col2:
        st.download_button("📥 Download report.json (Base Hypothesis)", safe_json_dumps(rep, indent=1), "report.json", mime="application/json", key="dl_report_json")
    if st.session_state.get("triage_files"):
        st.subheader("Expert Triage Files")
        t_cols = st.columns(3)
        for idx, (fn, content) in enumerate(st.session_state["triage_files"].items()):
            with t_cols[idx]:
                st.download_button(f"📥 Download {fn}", content, fn, key=f"dl_{fn}")

    # Continuous Follow-Up Section
    st.divider()
    st.subheader("💬 Continuous Follow-Up Research & Refinement")
    st.caption("Ask questions, test edge conditions, or refine the hypothesis. The engine maintains continuous research across turns: querying literature, tracking resolved vs. persisting flags, and generating formal follow-up reports. Type **exit** or click **🛑 End Session** in the sidebar to finish.")

    for idx, msg in enumerate(st.session_state.get("follow_up_history", [])):
        with st.chat_message(msg["role"]):
            if msg.get("report") and msg.get("turn"):
                f_rep = msg["report"]
                f_counts = f_rep.get("counts") or engine.count_flags(f_rep.get("active_issues", f_rep.get("issues", [])), f_rep.get("green_flags", []))
                fcol_r, fcol_y, fcol_g = st.columns(3)
                fcol_r.metric("🚩 Red Flags", f_counts.get("red", 0))
                fcol_y.metric("🟨 Yellow Flags", f_counts.get("yellow", 0))
                fcol_g.metric("🟩 Green Flags", f_counts.get("green", 0))
            st.markdown(msg["content"])
            if msg.get("report") and msg.get("turn"):
                turn_num = msg["turn"]
                f_rep = msg["report"]
                f_md = msg["content"]
                c1, c2 = st.columns(2)
                with c1:
                    st.download_button(
                        f"📥 Download report_followup_{turn_num}.md",
                        f_md,
                        f"report_followup_{turn_num}.md",
                        mime="text/markdown",
                        key=f"dl_f_{turn_num}_{idx}_md"
                    )
                with c2:
                    st.download_button(
                        f"📥 Download report_followup_{turn_num}.json",
                        safe_json_dumps(f_rep, indent=1),
                        f"report_followup_{turn_num}.json",
                        mime="application/json",
                        key=f"dl_f_{turn_num}_{idx}_json"
                    )

    if st.session_state.get("session_active", True):
        prompt = st.chat_input("Enter a follow-up prompt, probe, or refined hypothesis (or 'exit')...")
        if prompt:
            if prompt.strip().lower() in ("exit", "quit", "q", ":q"):
                st.session_state["session_active"] = False
                st.warning("Session ended. You can safely close this browser tab or restart a session.")
                st.rerun()
            else:
                st.session_state.setdefault("follow_up_history", []).append({"role": "user", "content": prompt})
                try:
                    gen = backend(gen_name)
                    crit = gen if critic_name == "same as generator" else backend(critic_name)
                    turn_idx = len(st.session_state.get("followup_reports", [])) + 1
                    with st.spinner(f"Conducting continuous literature research for follow-up turn {turn_idx}..."):
                        parent_to_use = st.session_state.get("latest_report") or rep
                        f_rep = engine.run_followup(
                            parent_to_use,
                            prompt,
                            books or books_default(corpus),
                            gen,
                            crit,
                            engine.LocalBM25(corpus),
                            turn_idx=turn_idx,
                            auto_entail=entail,
                            check_math=check_math,
                            log=lambda m: None
                        )
                    st.session_state.setdefault("followup_reports", []).append(f_rep)
                    st.session_state["latest_report"] = f_rep
                    f_md = engine.render_followup(f_rep)
                    st.session_state["follow_up_history"].append({
                        "role": "assistant",
                        "content": f_md,
                        "report": f_rep,
                        "turn": turn_idx
                    })
                    st.rerun()
                except Exception as e:
                    st.error(f"Error evaluating follow-up: {e}")

